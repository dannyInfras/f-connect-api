import os
import requests
import tempfile
from typing import Optional
from docx import Document
import fitz  # PyMuPDF
from logger import get_logger
from config import config

logger = get_logger('extractor')


def download_file(url: str) -> Optional[str]:
    """Download file from URL to temporary location."""
    try:
        response = requests.get(url, timeout=30)
        response.raise_for_status()
        
        # Extract file extension from URL
        url_path = url.split('?')[0]  # Remove query parameters
        file_ext = os.path.splitext(url_path)[1]
        
        # Create temporary file with correct extension
        with tempfile.NamedTemporaryFile(delete=False, suffix=file_ext) as temp_file:
            temp_file.write(response.content)
            temp_file_path = temp_file.name
        
        logger.info(f"Downloaded file from {url} to {temp_file_path}")
        return temp_file_path
        
    except Exception as e:
        logger.error(f"Error downloading file from {url}: {e}")
        return None


def extract_text_from_pdf(file_path: str) -> Optional[str]:
    """Extract text from PDF file using PyMuPDF."""
    try:
        text_content = []
        doc = fitz.open(file_path)
        
        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            text = page.get_text()
            if text.strip():
                text_content.append(text.strip())
        
        doc.close()
        
        if text_content:
            full_text = '\n\n'.join(text_content)
            logger.info(f"Extracted {len(full_text)} characters from PDF")
            return full_text
        else:
            logger.warning("No text content found in PDF")
            return None
            
    except Exception as e:
        logger.error(f"Error extracting text from PDF {file_path}: {e}")
        return None


def extract_text_from_docx(file_path: str) -> Optional[str]:
    """Extract text from DOCX file using python-docx."""
    try:
        doc = Document(file_path)
        text_content = []
        
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_content.append(paragraph.text.strip())
        
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text_content.append(' | '.join(row_text))
        
        if text_content:
            full_text = '\n'.join(text_content)
            logger.info(f"Extracted {len(full_text)} characters from DOCX")
            return full_text
        else:
            logger.warning("No text content found in DOCX")
            return None
            
    except Exception as e:
        logger.error(f"Error extracting text from DOCX {file_path}: {e}")
        return None


def extract_text_from_file(url: str) -> Optional[str]:
    """Extract text from file URL (supports PDF and DOCX)."""
    temp_file_path = None
    
    try:
        temp_file_path = download_file(url)
        if not temp_file_path:
            return None
        
        # Determine file type from original URL instead of temp file path
        url_path = url.split('?')[0]  # Remove query parameters
        file_ext = os.path.splitext(url_path)[1].lower()
        
        logger.info(f"Processing file with extension: {file_ext}")
        
        if file_ext == '.pdf':
            return extract_text_from_pdf(temp_file_path)
        elif file_ext in ['.docx', '.doc']:
            return extract_text_from_docx(temp_file_path)
        else:
            logger.error(f"Unsupported file type: {file_ext}")
            return None
            
    except Exception as e:
        logger.error(f"Error extracting text from file URL {url}: {e}")
        return None
        
    finally:
        if temp_file_path and os.path.exists(temp_file_path):
            try:
                os.unlink(temp_file_path)
                logger.debug(f"Cleaned up temporary file: {temp_file_path}")
            except Exception as e:
                logger.warning(f"Failed to clean up temporary file {temp_file_path}: {e}")


def clean_text_for_analysis(text: str) -> str:
    """Clean and prepare text for AI analysis."""
    if not text:
        return ""
    
    lines = [line.strip() for line in text.split('\n') if line.strip()]
    cleaned_text = '\n'.join(lines)
    
    # Use config value if available, otherwise default to 10000
    max_length = getattr(config, 'MAX_CV_TEXT_LENGTH', 10000)
    if len(cleaned_text) > max_length:
        cleaned_text = cleaned_text[:max_length] + "..."
        logger.info(f"Text truncated to {max_length} characters for analysis")
    
    return cleaned_text
